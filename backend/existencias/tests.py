import csv
import io

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from catalogos.models import (
    Cargo, Departamento, Direccion, Disco, EstadoImpresora, MarcaEquipo, MemoriaRAM,
    ModeloEquipo, Procesador, PulgadasMonitor, SistemaOperativo, TipoEquipo,
)
from config.seguridad import celda_segura
from usuarios.models import Tecnico
from .models import Equipo, Funcionario, ImpresoraRegistro, Monitor

# Pruebas sin bloqueo de login y con archivos estáticos simples (sin collectstatic).
PRUEBAS = dict(
    AXES_ENABLED=False,
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
)


class CeldaSeguraTests(TestCase):
    def test_neutraliza_formulas(self):
        for v in ["=1+1", "+1", "-1", "@SUMA(A1)", "\t=x"]:
            self.assertTrue(celda_segura(v).startswith("'"))

    def test_no_toca_texto_normal(self):
        self.assertEqual(celda_segura("JUAN PÉREZ"), "JUAN PÉREZ")
        self.assertEqual(celda_segura(5), 5)


@override_settings(**PRUEBAS)
class SoporteTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.dir = Direccion.objects.create(nombre="DITEC")
        cls.dep = Departamento.objects.create(direccion=cls.dir, nombre="Soporte")
        cls.cargo = Cargo.objects.create(nombre="Analista")
        cls.tipo = TipoEquipo.objects.create(nombre="Notebook")
        cls.marca = MarcaEquipo.objects.create(nombre="HP")
        cls.modelo = ModeloEquipo.objects.create(nombre="ProBook")
        cls.ram = MemoriaRAM.objects.create(tipo="8 GB")
        cls.disco = Disco.objects.create(tipo="SSD")
        cls.cpu = Procesador.objects.create(nombre="i5")
        cls.so = SistemaOperativo.objects.create(tipo="Windows 11")
        cls.pulgadas = PulgadasMonitor.objects.create(tipo="24")
        cls.estado = EstadoImpresora.objects.create(nombre="Operativa")
        cls.tecnico = Tecnico.objects.create_user("tecnico", password="ClaveTecnico-2026!")
        cls.jefe = Tecnico.objects.create_user("jefe", password="ClaveJefe-2026!", is_staff=True)
        func = Funcionario.objects.create(
            nombre='=HYPERLINK("http://malo","clic")', direccion=cls.dir,
            departamento=cls.dep, anexo="1234", cargo=cls.cargo,
        )
        cls.equipo = Equipo.objects.create(
            funcionario=func, tipo=cls.tipo, marca=cls.marca, modelo=cls.modelo,
            numero_serie="SN1", ram=cls.ram, disco=cls.disco, procesador=cls.cpu,
            sistema_operativo=cls.so, tecnico=cls.tecnico,
        )

    def _post_equipo(self, **extra):
        datos = {
            "funcionario_nombre": "Ana", "direccion": self.dir.pk, "departamento": self.dep.pk,
            "anexo": "", "cargo": self.cargo.pk, "tipo_equipo": self.tipo.pk,
            "marca": self.marca.pk, "modelo": self.modelo.pk, "numero_serie": "SN-NUEVO",
            "ram": self.ram.pk, "disco": self.disco.pk, "procesador": self.cpu.pk,
            "sistema_operativo": self.so.pk, "tiene_monitor": "no", "tiene_impresora": "no",
        }
        datos.update(extra)
        return self.client.post(reverse("ingresar_equipo"), datos)

    def test_anonimo_va_al_login(self):
        self.assertEqual(self.client.get(reverse("ver_equipos")).status_code, 302)

    def test_tecnico_no_puede_editar(self):
        self.client.force_login(self.tecnico)
        r = self.client.get(reverse("editar_equipo", args=[self.equipo.pk]))
        self.assertRedirects(r, reverse("ver_equipos"), fetch_redirect_response=False)

    def test_admin_puede_editar(self):
        self.client.force_login(self.jefe)
        r = self.client.get(reverse("editar_equipo", args=[self.equipo.pk]))
        self.assertEqual(r.status_code, 200)

    def test_registro_valido(self):
        self.client.force_login(self.tecnico)
        r = self._post_equipo()
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Equipo.objects.filter(numero_serie="SN-NUEVO").exists())

    def test_cantidad_de_monitores_basura_no_da_error_500(self):
        self.client.force_login(self.tecnico)
        r = self._post_equipo(tiene_monitor="si", cantidad_monitores="abc")
        self.assertEqual(r.status_code, 200)

    def test_cantidad_de_monitores_se_acota_a_dos(self):
        self.client.force_login(self.tecnico)
        datos = {"tiene_monitor": "si", "cantidad_monitores": "100000"}
        for i in (1, 2, 3):
            datos.update({f"monitor{i}_pulgadas": self.pulgadas.pk, f"monitor{i}_numero_serie": f"M{i}"})
        self._post_equipo(**datos)
        self.assertEqual(Monitor.objects.count(), 2)

    def test_id_no_numerico_no_da_error_500(self):
        self.client.force_login(self.tecnico)
        r = self._post_equipo(ram="xyz")
        self.assertEqual(r.status_code, 200)
        self.assertFalse(Equipo.objects.filter(numero_serie="SN-NUEVO").exists())

    def test_impresora_sin_estado_no_da_error_500(self):
        self.client.force_login(self.tecnico)
        r = self.client.post(reverse("ingresar_impresora"), {
            "direccion": self.dir.pk, "departamento": self.dep.pk, "numero_serie": "IMP1",
        })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(ImpresoraRegistro.objects.count(), 0)

    def test_excel_neutraliza_formulas(self):
        self.client.force_login(self.tecnico)
        r = self.client.get(reverse("exportar_excel"))
        valor = load_workbook(io.BytesIO(r.content)).active["A2"].value
        self.assertTrue(valor.startswith("'="))

    def test_csv_neutraliza_formulas(self):
        self.client.force_login(self.tecnico)
        r = self.client.get(reverse("exportar_csv"))
        filas = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        self.assertTrue(filas[1][0].startswith("'="), filas[:2])
        self.assertNotIn("\ufeff", r.content.decode("utf-8-sig"))  # BOM solo al inicio

    def test_excel_de_impresoras_funciona(self):
        ImpresoraRegistro.objects.create(
            direccion=self.dir, departamento=self.dep, estado=self.estado,
            numero_serie="IMP1", tecnico=self.tecnico,
        )
        self.client.force_login(self.tecnico)
        r = self.client.get(reverse("exportar_impresoras_excel"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(r.content)).active["E2"].value, "IMP1")
