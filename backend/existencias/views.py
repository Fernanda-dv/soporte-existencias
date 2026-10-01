import csv
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import transaction, IntegrityError, models
from django.http import HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import Funcionario, Equipo, Monitor, ImpresoraRegistro
from catalogos.models import (
    Direccion, Departamento, Cargo, TipoEquipo,
    MarcaEquipo, ModeloEquipo,
    MarcaMonitor, ModeloMonitor,
    MarcaImpresora, ModeloImpresora,
    MemoriaRAM, Disco, Procesador, SistemaOperativo, PulgadasMonitor, EstadoImpresora
)


from config.seguridad import celda_segura

# Errores que puede provocar un formulario manipulado (IDs que no existen o no numéricos):
# se informan al usuario en vez de producir un error 500.
ERRORES_DATOS = (IntegrityError, ValueError, ValidationError)
MAX_MONITORES = 2  # el formulario ofrece 1 o 2 pantallas


def _cantidad_monitores(valor):
    """Cantidad de monitores enviada por el navegador, acotada a 0..MAX_MONITORES."""
    try:
        return max(0, min(int(valor), MAX_MONITORES))
    except (TypeError, ValueError):
        return 0


def _fecha_local(fecha):
    """Las fechas se guardan en UTC; en las exportaciones se muestran en hora de Chile."""
    return timezone.localtime(fecha).strftime("%d/%m/%Y %H:%M")


def _fila_segura(valores):
    """Evita que un texto como =HYPERLINK(...) se ejecute como fórmula al abrir Excel/CSV."""
    return [celda_segura(v) for v in valores]


def es_admin(user):
    """Verifica si el usuario es superusuario o miembro del personal de administración."""
    return user.is_superuser or user.is_staff


# ==========================================
# 1. REGISTRO Y EDICIÓN DE EQUIPOS
# ==========================================

@login_required
@require_http_methods(["GET", "POST"])
def ingresar_equipo(request):
    if request.method == 'POST':
        # Datos Funcionario
        nombre_func = request.POST.get('funcionario_nombre', '').strip().upper()
        direccion_id = request.POST.get('direccion')
        departamento_id = request.POST.get('departamento')
        anexo_val = request.POST.get('anexo', '').strip() or None
        cargo_id = request.POST.get('cargo')

        # Datos Equipo (Marca y Modelo específicos de Equipo)
        tipo_equipo_id = request.POST.get('tipo_equipo')
        marca_id = request.POST.get('marca')
        modelo_id = request.POST.get('modelo')
        num_serie = request.POST.get('numero_serie', '').strip() or None
        num_inventario = request.POST.get('numero_inventario', '').strip() or None

        # Datos Periféricos
        tiene_monitor = request.POST.get('tiene_monitor') == 'si'
        cant_monitores = _cantidad_monitores(request.POST.get('cantidad_monitores')) if tiene_monitor else 0
        tiene_impresora = request.POST.get('tiene_impresora') == 'si'

        # Validaciones de Negocio
        errores = []

        if not marca_id or not modelo_id:
            errores.append("Debe seleccionar la Marca y el Modelo del equipo.")

        if not num_serie and not num_inventario:
            errores.append("Debe ingresar al menos el Número de Serie o el Número de Inventario del equipo.")

        if tiene_monitor:
            if cant_monitores < 1:
                errores.append("Debe indicar la cantidad de monitores.")
            for i in range(1, cant_monitores + 1):
                pulgadas = request.POST.get(f'monitor{i}_pulgadas')
                m_ns = request.POST.get(f'monitor{i}_numero_serie', '').strip()
                m_ni = request.POST.get(f'monitor{i}_numero_inventario', '').strip()
                if not pulgadas:
                    errores.append(f"Debe seleccionar las pulgadas para la Pantalla {i}.")
                if not m_ns and not m_ni:
                    errores.append(f"Para la Pantalla {i}, debe ingresar al menos el N° de Serie o el N° de Inventario.")

        if tiene_impresora:
            imp_ns = request.POST.get('impresora_numero_serie', '').strip()
            imp_ni = request.POST.get('impresora_numero_inventario', '').strip()
            if not imp_ns and not imp_ni:
                errores.append("Para la impresora del equipo, debe ingresar al menos el N° de Serie o el N° de Inventario.")

        if errores:
            for err in errores:
                messages.error(request, err)
        else:
            try:
                with transaction.atomic():
                    funcionario = Funcionario.objects.create(
                        nombre=nombre_func,
                        direccion_id=direccion_id,
                        departamento_id=departamento_id,
                        anexo=anexo_val,
                        cargo_id=cargo_id
                    )

                    equipo = Equipo.objects.create(
                        funcionario=funcionario,
                        tipo_id=tipo_equipo_id,
                        marca_id=marca_id,
                        modelo_id=modelo_id,
                        numero_serie=num_serie,
                        numero_inventario=num_inventario,
                        ram_id=request.POST.get('ram'),
                        disco_id=request.POST.get('disco'),
                        procesador_id=request.POST.get('procesador'),
                        sistema_operativo_id=request.POST.get('sistema_operativo'),
                        tiene_monitor=tiene_monitor,
                        tiene_impresora=tiene_impresora,
                        marca_impresora_id=request.POST.get('impresora_marca') or None if tiene_impresora else None,
                        modelo_impresora_id=request.POST.get('impresora_modelo') or None if tiene_impresora else None,
                        numero_serie_impresora=request.POST.get('impresora_numero_serie', '').strip() or None if tiene_impresora else None,
                        numero_inventario_impresora=request.POST.get('impresora_numero_inventario', '').strip() or None if tiene_impresora else None,
                        tecnico=request.user
                    )

                    if tiene_monitor:
                        for i in range(1, cant_monitores + 1):
                            Monitor.objects.create(
                                equipo=equipo,
                                pantalla_numero=i,
                                pulgadas_id=request.POST.get(f'monitor{i}_pulgadas'),
                                hdmi=request.POST.get(f'monitor{i}_hdmi') in ['True', 'on', 'true'],
                                marca_id=request.POST.get(f'monitor{i}_marca') or None,
                                modelo_id=request.POST.get(f'monitor{i}_modelo') or None,
                                numero_serie=request.POST.get(f'monitor{i}_numero_serie', '').strip() or None,
                                numero_inventario=request.POST.get(f'monitor{i}_numero_inventario', '').strip() or None,
                            )

                messages.success(request, "Equipo registrado correctamente.")
                return redirect('ver_equipos')
            except IntegrityError:
                messages.error(request, "El N° de Serie o N° de Inventario ya se encuentra registrado en el sistema.")
            except (ValueError, ValidationError):
                messages.error(request, "Hay datos inválidos en el formulario. Revise las listas seleccionadas.")

    context = {
        'direcciones': Direccion.objects.all(),
        'departamentos': Departamento.objects.all(),
        'cargos': Cargo.objects.all(),
        'tipos_equipo': TipoEquipo.objects.all(),
        'marcas_equipo': MarcaEquipo.objects.all(),
        'modelos_equipo': ModeloEquipo.objects.all(),
        'marcas_monitor': MarcaMonitor.objects.all(),
        'modelos_monitor': ModeloMonitor.objects.all(),
        'marcas_impresora': MarcaImpresora.objects.all(),
        'modelos_impresora': ModeloImpresora.objects.all(),
        'rams': MemoriaRAM.objects.all(),
        'discos': Disco.objects.all(),
        'procesadores': Procesador.objects.all(),
        'sistemas_operativos': SistemaOperativo.objects.all(),
        'pulgadas_monitor': PulgadasMonitor.objects.all(),
    }
    return render(request, 'form.html', context)


@login_required
@require_http_methods(["GET", "POST"])
def editar_equipo(request, equipo_id):
    if not es_admin(request.user):
        messages.error(request, "No tiene permisos para editar registros.")
        return redirect('ver_equipos')

    equipo = get_object_or_404(Equipo, id=equipo_id)
    funcionario = equipo.funcionario

    if request.method == 'POST':
        nombre_func = request.POST.get('funcionario_nombre', '').strip().upper()
        num_serie = request.POST.get('numero_serie', '').strip() or None
        num_inventario = request.POST.get('numero_inventario', '').strip() or None

        tiene_monitor = request.POST.get('tiene_monitor') == 'si'
        cant_monitores = _cantidad_monitores(request.POST.get('cantidad_monitores')) if tiene_monitor else 0
        tiene_impresora = request.POST.get('tiene_impresora') == 'si'

        errores = []

        if not request.POST.get('marca') or not request.POST.get('modelo'):
            errores.append("Debe seleccionar la Marca y el Modelo del equipo.")

        if not num_serie and not num_inventario:
            errores.append("Debe ingresar al menos el N° de Serie o N° de Inventario del equipo.")

        if tiene_monitor:
            for i in range(1, cant_monitores + 1):
                m_ns = request.POST.get(f'monitor{i}_numero_serie', '').strip()
                m_ni = request.POST.get(f'monitor{i}_numero_inventario', '').strip()
                if not m_ns and not m_ni:
                    errores.append(f"Para la Pantalla {i}, debe ingresar al menos el N° de Serie o el N° de Inventario.")

        if tiene_impresora:
            imp_ns = request.POST.get('impresora_numero_serie', '').strip()
            imp_ni = request.POST.get('impresora_numero_inventario', '').strip()
            if not imp_ns and not imp_ni:
                errores.append("Para la impresora del equipo, debe ingresar al menos el N° de Serie o N° de Inventario.")

        if errores:
            for err in errores:
                messages.error(request, err)
        else:
            try:
                with transaction.atomic():
                    funcionario.nombre = nombre_func
                    funcionario.direccion_id = request.POST.get('direccion')
                    funcionario.departamento_id = request.POST.get('departamento')
                    funcionario.anexo = request.POST.get('anexo', '').strip() or None
                    funcionario.cargo_id = request.POST.get('cargo')
                    funcionario.save()

                    equipo.tipo_id = request.POST.get('tipo_equipo')
                    equipo.marca_id = request.POST.get('marca')
                    equipo.modelo_id = request.POST.get('modelo')
                    equipo.numero_serie = num_serie
                    equipo.numero_inventario = num_inventario
                    equipo.ram_id = request.POST.get('ram')
                    equipo.disco_id = request.POST.get('disco')
                    equipo.procesador_id = request.POST.get('procesador')
                    equipo.sistema_operativo_id = request.POST.get('sistema_operativo')
                    equipo.tiene_monitor = tiene_monitor
                    equipo.tiene_impresora = tiene_impresora

                    equipo.marca_impresora_id = request.POST.get('impresora_marca') or None if tiene_impresora else None
                    equipo.modelo_impresora_id = request.POST.get('impresora_modelo') or None if tiene_impresora else None
                    equipo.numero_serie_impresora = request.POST.get('impresora_numero_serie', '').strip() or None if tiene_impresora else None
                    equipo.numero_inventario_impresora = request.POST.get('impresora_numero_inventario', '').strip() or None if tiene_impresora else None
                    equipo.save()

                    equipo.monitores.all().delete()
                    if tiene_monitor:
                        for i in range(1, cant_monitores + 1):
                            Monitor.objects.create(
                                equipo=equipo,
                                pantalla_numero=i,
                                pulgadas_id=request.POST.get(f'monitor{i}_pulgadas'),
                                hdmi=request.POST.get(f'monitor{i}_hdmi') in ['True', 'on', 'true'],
                                marca_id=request.POST.get(f'monitor{i}_marca') or None,
                                modelo_id=request.POST.get(f'monitor{i}_modelo') or None,
                                numero_serie=request.POST.get(f'monitor{i}_numero_serie', '').strip() or None,
                                numero_inventario=request.POST.get(f'monitor{i}_numero_inventario', '').strip() or None,
                            )

                messages.success(request, "Registro de equipo actualizado correctamente.")
                return redirect('ver_equipos')
            except IntegrityError:
                messages.error(request, "El N° de Serie o N° de Inventario ya pertenece a otro registro.")
            except (ValueError, ValidationError):
                messages.error(request, "Hay datos inválidos en el formulario. Revise las listas seleccionadas.")

    monitor1 = equipo.monitores.filter(pantalla_numero=1).first()
    monitor2 = equipo.monitores.filter(pantalla_numero=2).first()

    context = {
        'equipo': equipo,
        'funcionario': funcionario,
        'monitor1': monitor1,
        'monitor2': monitor2,
        'direcciones': Direccion.objects.all(),
        'departamentos': Departamento.objects.all(),
        'cargos': Cargo.objects.all(),
        'tipos_equipo': TipoEquipo.objects.all(),
        'marcas_equipo': MarcaEquipo.objects.all(),
        'modelos_equipo': ModeloEquipo.objects.all(),
        'marcas_monitor': MarcaMonitor.objects.all(),
        'modelos_monitor': ModeloMonitor.objects.all(),
        'marcas_impresora': MarcaImpresora.objects.all(),
        'modelos_impresora': ModeloImpresora.objects.all(),
        'rams': MemoriaRAM.objects.all(),
        'discos': Disco.objects.all(),
        'procesadores': Procesador.objects.all(),
        'sistemas_operativos': SistemaOperativo.objects.all(),
        'pulgadas_monitor': PulgadasMonitor.objects.all(),
    }
    return render(request, 'form.html', context)


# ==========================================
# 2. REGISTRO DE IMPRESORAS INDEPENDIENTES
# ==========================================

@login_required
@require_http_methods(["GET", "POST"])
def ingresar_impresora(request):
    if request.method == 'POST':
        direccion_id = request.POST.get('direccion')
        departamento_id = request.POST.get('departamento')
        marca_id = request.POST.get('marca')
        modelo_id = request.POST.get('modelo')
        ns = request.POST.get('numero_serie', '').strip() or None
        ni = request.POST.get('numero_inventario', '').strip() or None
        estado_id = request.POST.get('estado')

        if not ns and not ni:
            messages.error(request, "Debe ingresar al menos el N° de Serie o el N° de Inventario de la impresora.")
        else:
            try:
                with transaction.atomic():
                    ImpresoraRegistro.objects.create(
                        direccion_id=direccion_id,
                        departamento_id=departamento_id,
                        marca_id=marca_id or None,
                        modelo_id=modelo_id or None,
                        numero_serie=ns,
                        numero_inventario=ni,
                        estado_id=estado_id,
                        tecnico=request.user
                    )
                messages.success(request, "Impresora registrada exitosamente.")
                return redirect('ver_equipos')
            except ERRORES_DATOS:
                messages.error(request, "Faltan datos obligatorios (Dirección, Departamento o Estado) o hay datos inválidos.")

    # Contexto específico para form_impresora.html con sólo MarcaImpresora y ModeloImpresora
    context = {
        'direcciones': Direccion.objects.all(),
        'departamentos': Departamento.objects.all(),
        'marcas_impresora': MarcaImpresora.objects.all(),
        'modelos_impresora': ModeloImpresora.objects.all(),
        'estados_impresora': EstadoImpresora.objects.all(),
    }
    return render(request, 'form_impresora.html', context)


# ==========================================
# 3. TABLA DE INVENTARIO CON FILTROS SEPARADOS
# ==========================================

@login_required
def ver_equipos(request):
    q_eq = request.GET.get('q_eq', request.GET.get('q', '')).strip()
    tecnico_eq = request.GET.get('tecnico_eq', request.GET.get('tecnico', 'todos'))

    q_imp = request.GET.get('q_imp', '').strip()
    tecnico_imp = request.GET.get('tecnico_imp', 'todos')

    equipos = Equipo.objects.select_related(
        'funcionario', 'funcionario__direccion', 'funcionario__departamento', 'funcionario__cargo',
        'tipo', 'marca', 'modelo', 'ram', 'disco', 'procesador', 'sistema_operativo',
        'marca_impresora', 'modelo_impresora', 'tecnico'
    ).prefetch_related('monitores', 'monitores__marca', 'monitores__modelo', 'monitores__pulgadas').order_by('-fecha_creacion')

    impresoras = ImpresoraRegistro.objects.select_related(
        'direccion', 'departamento', 'marca', 'modelo', 'estado', 'tecnico'
    ).order_by('-fecha_creacion')

    if tecnico_eq == 'mio':
        equipos = equipos.filter(tecnico=request.user)

    if q_eq:
        equipos = equipos.filter(
            models.Q(funcionario__nombre__icontains=q_eq) |
            models.Q(funcionario__direccion__nombre__icontains=q_eq) |
            models.Q(funcionario__departamento__nombre__icontains=q_eq) |
            models.Q(funcionario__anexo__icontains=q_eq) |
            models.Q(tipo__nombre__icontains=q_eq) |
            models.Q(marca__nombre__icontains=q_eq) |
            models.Q(modelo__nombre__icontains=q_eq) |
            models.Q(numero_serie__icontains=q_eq) |
            models.Q(numero_inventario__icontains=q_eq)
        ).distinct()

    if tecnico_imp == 'mio':
        impresoras = impresoras.filter(tecnico=request.user)

    if q_imp:
        impresoras = impresoras.filter(
            models.Q(direccion__nombre__icontains=q_imp) |
            models.Q(departamento__nombre__icontains=q_imp) |
            models.Q(marca__nombre__icontains=q_imp) |
            models.Q(modelo__nombre__icontains=q_imp) |
            models.Q(numero_serie__icontains=q_imp) |
            models.Q(numero_inventario__icontains=q_imp)
        ).distinct()

    context = {
        'equipos': equipos,
        'impresoras': impresoras,
        'q_eq': q_eq,
        'tecnico_eq': tecnico_eq,
        'q_imp': q_imp,
        'tecnico_imp': tecnico_imp,
        'es_admin': es_admin(request.user),
    }
    return render(request, 'tabla.html', context)


# ==========================================
# 4. VISTAS DE EXPORTACIÓN (EXCEL / CSV)
# ==========================================

@login_required
def exportar_excel(request):
    tecnico_eq = request.GET.get('tecnico', 'todos')
    q_eq = request.GET.get('q', '').strip()

    equipos = Equipo.objects.select_related(
        'funcionario', 'funcionario__direccion', 'funcionario__departamento', 'funcionario__cargo',
        'tipo', 'marca', 'modelo', 'ram', 'disco', 'procesador', 'sistema_operativo',
        'marca_impresora', 'modelo_impresora', 'tecnico'
    ).prefetch_related('monitores', 'monitores__marca', 'monitores__modelo', 'monitores__pulgadas').order_by('-fecha_creacion')

    if tecnico_eq == 'mio':
        equipos = equipos.filter(tecnico=request.user)

    if q_eq:
        equipos = equipos.filter(
            models.Q(funcionario__nombre__icontains=q_eq) |
            models.Q(numero_serie__icontains=q_eq) |
            models.Q(numero_inventario__icontains=q_eq)
        ).distinct()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Equipos"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='10174A', end_color='10174A', fill_type='solid')

    headers = [
        "Funcionario", "Dirección", "Departamento", "Anexo", "Cargo",
        "Tipo Equipo", "Marca Equipo", "Modelo Equipo", "N° Serie Equipo", "N° Inventario Equipo",
        "RAM", "Disco", "Procesador", "S.O.",
        "Detalle Monitores (Marca / Modelo / N° Serie / N° Inv)",
        "Detalle Impresora (Marca / Modelo / N° Serie / N° Inv)",
        "Fecha Creación", "Técnico Registrador"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for eq in equipos:
        # Formatear detalle de monitores
        mon_str = []
        if eq.tiene_monitor:
            for m in eq.monitores.all():
                m_marca = m.marca.nombre if m.marca else "S/M"
                m_mod = m.modelo.nombre if m.modelo else "S/M"
                m_ns = m.numero_serie or "N/A"
                m_ni = m.numero_inventario or "N/A"
                mon_str.append(f"P{m.pantalla_numero}: [{m_marca} {m_mod}] Serie: {m_ns} | Inv: {m_ni}")

        # Formatear detalle de impresora
        imp_str = "No"
        if eq.tiene_impresora:
            imp_marca = eq.marca_impresora.nombre if eq.marca_impresora else "S/M"
            imp_mod = eq.modelo_impresora.nombre if eq.modelo_impresora else "S/M"
            imp_ns = eq.numero_serie_impresora or "N/A"
            imp_ni = eq.numero_inventario_impresora or "N/A"
            imp_str = f"[{imp_marca} {imp_mod}] Serie: {imp_ns} | Inv: {imp_ni}"

        ws.append(_fila_segura([
            eq.funcionario.nombre,
            eq.funcionario.direccion.nombre,
            eq.funcionario.departamento.nombre,
            eq.funcionario.anexo or "N/A",
            eq.funcionario.cargo.nombre,
            eq.tipo.nombre,
            eq.marca.nombre if eq.marca else "N/A",
            eq.modelo.nombre if eq.modelo else "N/A",
            eq.numero_serie or "N/A",
            eq.numero_inventario or "N/A",
            eq.ram.tipo,
            eq.disco.tipo,
            eq.procesador.nombre,
            eq.sistema_operativo.tipo,
            " // ".join(mon_str) if mon_str else "No",
            imp_str,
            _fecha_local(eq.fecha_creacion),
            eq.tecnico.get_full_name() or eq.tecnico.username
        ]))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventario_equipos_maipu.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar_csv(request):
    tecnico_eq = request.GET.get('tecnico', 'todos')
    q_eq = request.GET.get('q', '').strip()

    equipos = Equipo.objects.select_related(
        'funcionario', 'funcionario__direccion', 'funcionario__departamento', 'funcionario__cargo',
        'tipo', 'marca', 'modelo', 'ram', 'disco', 'procesador', 'sistema_operativo',
        'marca_impresora', 'modelo_impresora', 'tecnico'
    ).prefetch_related('monitores', 'monitores__marca', 'monitores__modelo', 'monitores__pulgadas').order_by('-fecha_creacion')

    if tecnico_eq == 'mio':
        equipos = equipos.filter(tecnico=request.user)

    if q_eq:
        equipos = equipos.filter(
            models.Q(funcionario__nombre__icontains=q_eq) |
            models.Q(numero_serie__icontains=q_eq) |
            models.Q(numero_inventario__icontains=q_eq)
        ).distinct()

    # utf-8 + BOM escrito UNA vez: con charset=utf-8-sig Django repetía el BOM al inicio de cada fila.
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response.write('\ufeff')
    response['Content-Disposition'] = 'attachment; filename="inventario_equipos_maipu.csv"'
    writer = csv.writer(response, delimiter=';')

    writer.writerow([
        "Funcionario", "Dirección", "Departamento", "Anexo", "Cargo",
        "Tipo Equipo", "Marca Equipo", "Modelo Equipo", "N° Serie Equipo", "N° Inventario Equipo",
        "RAM", "Disco", "Procesador", "S.O.",
        "Detalle Monitores (Marca / Modelo / N° Serie / N° Inv)",
        "Detalle Impresora (Marca / Modelo / N° Serie / N° Inv)",
        "Fecha Creación", "Técnico Registrador"
    ])

    for eq in equipos:
        mon_str = []
        if eq.tiene_monitor:
            for m in eq.monitores.all():
                m_marca = m.marca.nombre if m.marca else "S/M"
                m_mod = m.modelo.nombre if m.modelo else "S/M"
                m_ns = m.numero_serie or "N/A"
                m_ni = m.numero_inventario or "N/A"
                mon_str.append(f"P{m.pantalla_numero}: [{m_marca} {m_mod}] Serie: {m_ns} | Inv: {m_ni}")

        imp_str = "No"
        if eq.tiene_impresora:
            imp_marca = eq.marca_impresora.nombre if eq.marca_impresora else "S/M"
            imp_mod = eq.modelo_impresora.nombre if eq.modelo_impresora else "S/M"
            imp_ns = eq.numero_serie_impresora or "N/A"
            imp_ni = eq.numero_inventario_impresora or "N/A"
            imp_str = f"[{imp_marca} {imp_mod}] Serie: {imp_ns} | Inv: {imp_ni}"

        writer.writerow(_fila_segura([
            eq.funcionario.nombre,
            eq.funcionario.direccion.nombre,
            eq.funcionario.departamento.nombre,
            eq.funcionario.anexo or "N/A",
            eq.funcionario.cargo.nombre,
            eq.tipo.nombre,
            eq.marca.nombre if eq.marca else "N/A",
            eq.modelo.nombre if eq.modelo else "N/A",
            eq.numero_serie or "N/A",
            eq.numero_inventario or "N/A",
            eq.ram.tipo,
            eq.disco.tipo,
            eq.procesador.nombre,
            eq.sistema_operativo.tipo,
            " // ".join(mon_str) if mon_str else "No",
            imp_str,
            _fecha_local(eq.fecha_creacion),
            eq.tecnico.get_full_name() or eq.tecnico.username
        ]))

    return response

@login_required
def exportar_impresoras_excel(request):
    tecnico_imp = request.GET.get('tecnico', 'todos')
    q_imp = request.GET.get('q', '').strip()

    impresoras = ImpresoraRegistro.objects.select_related(
        'direccion', 'departamento', 'marca', 'modelo', 'estado', 'tecnico'
    ).order_by('-fecha_creacion')

    if tecnico_imp == 'mio':
        impresoras = impresoras.filter(tecnico=request.user)

    if q_imp:
        impresoras = impresoras.filter(
            models.Q(direccion__nombre__icontains=q_imp) |
            models.Q(departamento__nombre__icontains=q_imp) |
            models.Q(marca__nombre__icontains=q_imp) |
            models.Q(modelo__nombre__icontains=q_imp) |
            models.Q(numero_serie__icontains=q_imp) |
            models.Q(numero_inventario__icontains=q_imp)
        ).distinct()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Impresoras"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='10174A', end_color='10174A', fill_type='solid')

    headers = ["Dirección", "Departamento", "Marca", "Modelo", "N° Serie", "N° Inventario", "Estado", "Fecha Creación", "Técnico Registrador"]
    ws.append(headers)

    for cell in ws[1]:  # fila de encabezados (antes recorría filas y daba error 500)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for imp in impresoras:
        ws.append(_fila_segura([
            imp.direccion.nombre if imp.direccion else "N/A",
            imp.departamento.nombre if imp.departamento else "N/A",
            imp.marca.nombre if imp.marca else "N/A",
            imp.modelo.nombre if imp.modelo else "N/A",
            imp.numero_serie or "N/A",
            imp.numero_inventario or "N/A",
            imp.estado.nombre if imp.estado else "N/A",
            _fecha_local(imp.fecha_creacion),
            imp.tecnico.get_full_name() or imp.tecnico.username
        ]))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventario_impresoras_maipu.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar_impresoras_csv(request):
    tecnico_imp = request.GET.get('tecnico', 'todos')
    q_imp = request.GET.get('q', '').strip()

    impresoras = ImpresoraRegistro.objects.select_related(
        'direccion', 'departamento', 'marca', 'modelo', 'estado', 'tecnico'
    ).order_by('-fecha_creacion')

    if tecnico_imp == 'mio':
        impresoras = impresoras.filter(tecnico=request.user)

    if q_imp:
        impresoras = impresoras.filter(
            models.Q(direccion__nombre__icontains=q_imp) |
            models.Q(departamento__nombre__icontains=q_imp) |
            models.Q(marca__nombre__icontains=q_imp) |
            models.Q(modelo__nombre__icontains=q_imp) |
            models.Q(numero_serie__icontains=q_imp) |
            models.Q(numero_inventario__icontains=q_imp)
        ).distinct()

    # utf-8 + BOM escrito UNA vez: con charset=utf-8-sig Django repetía el BOM al inicio de cada fila.
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response.write('\ufeff')
    response['Content-Disposition'] = 'attachment; filename="inventario_impresoras_maipu.csv"'
    writer = csv.writer(response, delimiter=';')

    writer.writerow(["Dirección", "Departamento", "Marca", "Modelo", "N° Serie", "N° Inventario", "Estado", "Fecha Creación", "Técnico Registrador"])

    for imp in impresoras:
        writer.writerow(_fila_segura([
            imp.direccion.nombre if imp.direccion else "N/A",
            imp.departamento.nombre if imp.departamento else "N/A",
            imp.marca.nombre if imp.marca else "N/A",
            imp.modelo.nombre if imp.modelo else "N/A",
            imp.numero_serie or "N/A",
            imp.numero_inventario or "N/A",
            imp.estado.nombre if imp.estado else "N/A",
            _fecha_local(imp.fecha_creacion),
            imp.tecnico.get_full_name() or imp.tecnico.username
        ]))

    return response